import html as html_lib
import io
from datetime import date

import openpyxl
from openpyxl.styles import Font

from app.models import Rfi, VeredictoRfi

_VEREDITO_LABEL = {
    VeredictoRfi.atende: "Atende",
    VeredictoRfi.atende_parcial: "Atende parcial",
    VeredictoRfi.nao_atende: "Não atende",
    VeredictoRfi.sem_evidencia: "Sem evidência",
}

_VEREDITO_COLOR = {
    VeredictoRfi.atende: "#1f9a5e",
    VeredictoRfi.atende_parcial: "#92660a",
    VeredictoRfi.nao_atende: "#c2412f",
    VeredictoRfi.sem_evidencia: "#56706a",
}


def _esc(valor: str | None) -> str:
    return html_lib.escape(valor or "")


def _itens_por_categoria(rfi: Rfi) -> dict[str, list]:
    por_categoria: dict[str, list] = {}
    for item in sorted(rfi.itens, key=lambda i: i.ordem):
        por_categoria.setdefault(item.categoria or "Sem categoria", []).append(item)
    return por_categoria


def _fontes_texto(item) -> str:
    return "; ".join(f.get("titulo") or f.get("url", "") for f in (item.fontes or []))


def montar_rfi_export_html(rfi: Rfi) -> str:
    total = len(rfi.itens)
    avaliados = sum(1 for i in rfi.itens if i.veredito is not None)

    linhas = []
    for categoria, itens in _itens_por_categoria(rfi).items():
        linhas.append(f'<tr class="categoria-row"><td colspan="5"><strong>{_esc(categoria)}</strong></td></tr>')
        for item in itens:
            if item.veredito is not None:
                cor = _VEREDITO_COLOR[item.veredito]
                veredito_html = f'<span style="color:{cor};font-weight:600">{_VEREDITO_LABEL[item.veredito]}</span>'
            else:
                veredito_html = '<span style="color:#888">Em processamento…</span>'
            fontes_html = "".join(
                f'<div><a href="{_esc(f.get("url",""))}" target="_blank" rel="noreferrer">{_esc(f.get("titulo") or f.get("url",""))}</a></div>'
                for f in (item.fontes or [])
            )
            linhas.append(
                f"<tr><td>{_esc(item.codigo)}</td><td>{_esc(item.descricao)}</td>"
                f"<td>{veredito_html}</td><td>{_esc(item.justificativa)}</td><td>{fontes_html}</td></tr>"
            )

    return f"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<title>{_esc(rfi.titulo)} — Análise de RFI</title>
<style>
  body {{ font-family: Arial, Helvetica, sans-serif; color: #1a1a1a; max-width: 1100px; margin: 40px auto; padding: 0 20px; }}
  h1 {{ margin-bottom: 4px; }}
  .sub {{ color: #555; margin-bottom: 20px; }}
  table {{ width: 100%; border-collapse: collapse; margin-bottom: 8px; }}
  th, td {{ border: 1px solid #ccc; padding: 6px 10px; text-align: left; font-size: 13px; vertical-align: top; }}
  th {{ background: #f0f0f0; }}
  .categoria-row td {{ background: #fafafa; }}
  .pct {{ font-size: 20px; font-weight: 700; margin: 0 0 20px; }}
  @media print {{ body {{ margin: 0; }} }}
</style>
</head>
<body>
  <h1>{_esc(rfi.titulo)}</h1>
  <div class="sub">
    {_esc(rfi.orgao)}{f" · processo {_esc(rfi.numero_processo)}" if rfi.numero_processo else ""}
    · gerado em {date.today().strftime('%d/%m/%Y')}
  </div>
  <p class="pct">{avaliados}/{total} itens avaliados</p>

  <table>
    <thead><tr><th>Item</th><th>Descrição</th><th>Veredito</th><th>Justificativa</th><th>Fonte(s)</th></tr></thead>
    <tbody>{"".join(linhas)}</tbody>
  </table>
</body>
</html>"""


def montar_rfi_export_xlsx(rfi: Rfi) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RFI"

    cabecalho = ["Item", "Categoria", "Descrição", "Veredito", "Justificativa", "Fonte(s)"]
    ws.append(cabecalho)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for item in sorted(rfi.itens, key=lambda i: i.ordem):
        veredito = _VEREDITO_LABEL[item.veredito] if item.veredito is not None else "Em processamento"
        ws.append(
            [
                item.codigo,
                item.categoria or "",
                item.descricao,
                veredito,
                item.justificativa or "",
                _fontes_texto(item),
            ]
        )

    larguras = [12, 28, 60, 16, 50, 40]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
